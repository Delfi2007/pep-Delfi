"""Turn the downloaded government / official datasets into one JSON for the dashboard.

Inputs (data/gov/, downloaded by download_data.py):
  ncrb_cyber_crimes_against_children_2017_2021.xlsx
      NCRB "Crime in India" Table 9A.11 (state-wise cyber crimes against children, 2017-2021)
      plus 2011-census child population, compiled by V. Tikhute on Mendeley Data
      (doi:10.17632/mc7wsp8v9y.2, CC BY 4.0). Original source: ncrb.gov.in (public domain).
  ncmec_<year>_reports_by_country.pdf   (2019-2024)
      NCMEC CyberTipline reports by country.

Output: data/gov/gov_stats.json

Also fits two small, clearly-labelled analyses on the government data:
  * a log-linear growth trend of national cyber crimes against children, and
  * k-means clustering of states by per-capita rate and crime-type mix (2021).
"""
import json
import re
from pathlib import Path

import numpy as np
import openpyxl
from pypdf import PdfReader
from sklearn.cluster import KMeans
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parent.parent
GOV = ROOT / "data" / "gov"
XLSX = GOV / "ncrb_cyber_crimes_against_children_2017_2021.xlsx"

HEADS = ["cyber_blackmail_threat_harassment", "fake_profile", "cyber_pornography_children",
         "cyber_stalking_bullying", "online_games", "other", "total"]
HEAD_LABELS = {
    "cyber_blackmail_threat_harassment": "Cyber blackmail / threats / harassment",
    "fake_profile": "Fake profile",
    "cyber_pornography_children": "Publishing sexual material depicting children (IT Act 67B)",
    "cyber_stalking_bullying": "Cyber stalking / bullying",
    "online_games": "Crimes through online games",
    "other": "Other cyber crimes against children",
    "total": "Total",
}

# Name harmonisation across years (UT reorganisation in 2019/2020).
ALIASES = {"D&N Haveli and Daman & Diu": "D&N Haveli and Daman & Diu",
           "Delhi UT": "Delhi", "Jammu & Kashmir*": "Jammu & Kashmir",
           "D&N Haveli+": "D&N Haveli and Daman & Diu", "Daman & Diu+": "D&N Haveli and Daman & Diu"}

# Published national totals for later years that are not in the Mendeley workbook.
# Source: NCRB Crime in India 2022, as reported by Deccan Herald (Dec 2023).
REPORTED_LATER = [{"year": 2022, "total": 1823, "source": "NCRB Crime in India 2022 (as reported)",
                   "url": "https://deccanherald.com/india/cybercrimes-against-children-see-32-rise-in-a-year-ncrb-report-shows-2799597"}]


def _num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 0


def parse_ncrb():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    population = {}
    for row in wb["Child population"].iter_rows(min_row=2, values_only=True):
        if row[1] and isinstance(row[2], (int, float)):
            population[str(row[1]).strip()] = float(row[2])

    years = {}
    for ws in wb.worksheets:
        m = re.match(r"CII (\d{4})", ws.title)
        if not m:
            continue
        year = int(m.group(1))
        states, national, kind = [], None, "state"
        for row in ws.iter_rows(values_only=True):
            first = str(row[0]).strip() if row[0] is not None else ""
            label = first or (str(row[1]).strip() if row[1] is not None else "")
            if label.upper().startswith("UNION TERRITOR"):
                kind = "ut"
            if label.upper() == "TOTAL ALL INDIA":
                vals = [v for v in row[1:] if isinstance(v, (int, float))]
                national = dict(zip(HEADS, [_num(v) for v in vals[-7:]]))
                continue
            if not first.isdigit() or row[1] is None:
                continue
            name = ALIASES.get(str(row[1]).strip(), str(row[1]).strip())
            rec = {"state": name, "kind": kind, **dict(zip(HEADS, [_num(v) for v in row[2:9]]))}
            states.append(rec)
        years[year] = {"national": national, "states": states}
    return population, years


def ncrb_analysis(population, years):
    ys = sorted(years)
    national = [{"year": y, **years[y]["national"]} for y in ys]

    # growth trend: log-linear fit on 2017-2021 official values (+ 2022 reported value)
    pts = [(n["year"], n["total"]) for n in national] + [(r["year"], r["total"]) for r in REPORTED_LATER]
    x = np.array([p[0] for p in pts], dtype=float)
    y = np.log(np.array([p[1] for p in pts], dtype=float))
    slope, intercept = np.polyfit(x, y, 1)
    fitted = [{"year": int(v), "fitted": round(float(np.exp(intercept + slope * v)), 1)} for v in range(int(x.min()), int(x.max()) + 1)]
    trend = {"annual_growth_pct": round((np.exp(slope) - 1) * 100, 1), "points": fitted,
             "method": "Least-squares fit of log(total) on year (6 data points). Descriptive only — "
                       "growth partly reflects better reporting, not just more offending."}

    # per-capita and composition, latest year
    latest = max(ys)
    rows = []
    for s in years[latest]["states"]:
        pop = population.get(s["state"])
        rate = round(s["total"] / pop, 3) if pop else None   # per lakh children
        rows.append({**s, "child_population_lakh": pop, "rate_per_lakh_children": rate})

    # state series (total per year)
    series = {}
    for y in ys:
        for s in years[y]["states"]:
            series.setdefault(s["state"], {})[str(y)] = s["total"]

    # k-means on 2021: rate + share of each crime head
    usable = [r for r in rows if r["rate_per_lakh_children"] is not None and r["total"] >= 5]
    feats = []
    for r in usable:
        tot = r["total"]
        feats.append([np.log1p(r["rate_per_lakh_children"])] +
                     [r[h] / tot for h in HEADS[:-1]])
    X = StandardScaler().fit_transform(np.array(feats))
    km = KMeans(n_clusters=3, n_init=20, random_state=42).fit(X)
    clusters = []
    for k in range(3):
        members = [usable[i] for i in range(len(usable)) if km.labels_[i] == k]
        mean_rate = float(np.mean([m["rate_per_lakh_children"] for m in members]))
        mix = {h: round(float(np.mean([m[h] / m["total"] for m in members])), 3) for h in HEADS[:-1]}
        dominant = max(mix, key=mix.get)
        clusters.append({"cluster": k, "states": sorted(m["state"] for m in members),
                         "mean_rate_per_lakh_children": round(mean_rate, 3),
                         "dominant_crime_head": HEAD_LABELS[dominant], "mix": mix})
    clusters.sort(key=lambda c: -c["mean_rate_per_lakh_children"])
    name_for = {c["cluster"]: i for i, c in enumerate(clusters)}
    for i, r in enumerate(usable):
        r["cluster"] = name_for[int(km.labels_[i])]
    for i, c in enumerate(clusters):
        c["cluster"] = i

    kerala = [{"year": y, **next((s for s in years[y]["states"] if s["state"] == "Kerala"), {})} for y in ys]
    return {
        "source": {
            "name": "NCRB Crime in India, Table 9A.11 — Cyber Crimes against Children (State/UT-wise)",
            "compiled_by": "V. Tikhute, Mendeley Data, doi:10.17632/mc7wsp8v9y.2 (CC BY 4.0)",
            "url": "https://data.mendeley.com/datasets/mc7wsp8v9y/2",
            "original": "https://ncrb.gov.in",
            "caveat": "As provided by States/UTs; NCRB notes States/UTs should not be compared purely on crime figures.",
        },
        "heads": HEAD_LABELS,
        "national": national,
        "reported_later": REPORTED_LATER,
        "trend": trend,
        "latest_year": latest,
        "states_latest": sorted(rows, key=lambda r: -r["total"]),
        "state_series": series,
        "kerala": kerala,
        "clusters": {"k": 3, "year": latest, "features": "log rate per lakh children + share of each crime head",
                     "items": clusters},
    }


COMPARE = ["India", "Philippines", "Pakistan", "Bangladesh", "Indonesia", "Vietnam", "Brazil",
           "United States", "Iraq", "Algeria"]


def parse_ncmec():
    out = {}
    for pdf in sorted(GOV.glob("ncmec_*_reports_by_country.pdf")):
        year = int(re.search(r"(\d{4})", pdf.name).group(1))
        text = re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in PdfReader(pdf).pages))
        row = {}
        for c in COMPARE:
            m = re.search(rf"(?<![A-Za-z]){re.escape(c)} ((?:\d[\d,]* ?){{1,3}})(?= [A-Z]|$)", text)
            if m:
                nums = [int(n.replace(",", "")) for n in m.group(1).split()]
                row[c] = nums[-1]           # 2024 lists referrals / informational / total
        tot = re.search(r"([\d.]+) million (?:reports|CyberTipline)", text)
        out[year] = {"countries": row, "global_million": float(tot.group(1)) if tot else None}
    return out


# Worldwide totals stated by NCMEC (only years we could confirm; others left blank rather than guessed).
GLOBAL_TOTALS_MILLION = {2019: 16.9, 2021: 29.3, 2022: 32.0, 2023: 36.2}


def main():
    population, years = parse_ncrb()
    ncrb = ncrb_analysis(population, years)
    ncmec = parse_ncmec()
    for y, d in ncmec.items():
        if d["global_million"] is None:
            d["global_million"] = GLOBAL_TOTALS_MILLION.get(y)
    ncmec_out = {
        "source": {"name": "NCMEC CyberTipline Reports by Country",
                   "url": "https://www.missingkids.org/gethelpnow/cybertipline/cybertiplinedata",
                   "caveat": "Counts reports whose geographic indicators resolve to a country; most come from "
                             "US platforms' mandatory reporting. 2024 figures are lower partly because NCMEC "
                             "began bundling duplicate reports."},
        "years": [{"year": y, **ncmec[y]} for y in sorted(ncmec)],
    }
    out = {"ncrb": ncrb, "ncmec": ncmec_out}
    (GOV / "gov_stats.json").write_text(json.dumps(out, indent=2))
    print("NCRB national:", [(n["year"], n["total"]) for n in ncrb["national"]])
    print("trend growth %/yr:", ncrb["trend"]["annual_growth_pct"])
    print("Kerala:", [(k["year"], k.get("total")) for k in ncrb["kerala"]])
    print("clusters:", [(c["mean_rate_per_lakh_children"], len(c["states"]), c["dominant_crime_head"]) for c in ncrb["clusters"]["items"]])
    print("NCMEC India:", [(y["year"], y["countries"].get("India"), y["global_million"]) for y in ncmec_out["years"]])
    print("saved", GOV / "gov_stats.json")


if __name__ == "__main__":
    main()
