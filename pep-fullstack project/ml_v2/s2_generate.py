"""Stage 2 — generate the synthetic multimodal dataset (non-explicit).

For each pool (A = training, B = testing) it writes:
  chats_{P}.jsonl           plain chat conversations
  sms/{P}/*.xml             SMS backups (SMS Backup & Restore format)  + sms_{P}.jsonl
  email/{P}/*.mbox          email threads                              + email_{P}.jsonl
  docs/{P}/*.pdf|*.docx     chat transcripts as PDF and Word           + docs_{P}.jsonl
  screens/{P}/*.jpg         phone-style chat screenshots               + screens_{P}.jsonl
  voice/{P}/<id>/*.wav      one voice note per message (offline TTS)   + voice_{P}.jsonl
  video/{P}/*.mp4           screen recording with voice-over           + video_{P}.jsonl

Every record keeps its ground truth (label, predator, messages) so detection
can be scored per input type after the media are converted back to text.
"""
import json
import subprocess
import textwrap
import time
import wave
from datetime import datetime, timedelta
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from xml.sax.saxutils import quoteattr

from PIL import Image, ImageDraw, ImageFont

from config import N, SYN, log, rng
from synth_text import conversation

FONT = "C:/Windows/Fonts/arial.ttf"
FONT_B = "C:/Windows/Fonts/arialbd.ttf"
W, H = 720, 1280


def write_jsonl(path, rows):
    with open(path, "w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def make(pool, n, salt, modality, kind_mix=None, max_beats=None):
    r = rng(salt + (0 if pool == "A" else 50_000))
    rows = []
    for i in range(n):
        kind = None
        if kind_mix == "balanced":
            kind = "groom" if i % 2 == 0 else ("hard" if i % 4 == 1 else "benign")
        c = conversation(pool, r, kind=kind, max_beats=max_beats)
        c.update({"id": f"{modality}-{pool}-{i:05d}", "pool": pool, "modality": modality})
        rows.append(c)
    return rows


def timestamps(r, n, start=None):
    t = start or datetime(2026, r.randint(1, 9), r.randint(1, 28), r.randint(7, 22), r.randint(0, 59))
    out = []
    for _ in range(n):
        t += timedelta(seconds=r.randint(20, 900))
        out.append(t)
    return out


# ---------------------------------------------------------------- chats
def gen_chats(pool):
    rows = make(pool, N["chats"], 1, "chat")
    write_jsonl(SYN / f"chats_{pool}.jsonl", rows)
    return len(rows)


# ---------------------------------------------------------------- SMS
def gen_sms(pool):
    rows = make(pool, N["sms_threads"], 2, "sms", max_beats=5)
    r = rng(22)
    out = SYN / "sms" / pool
    out.mkdir(parents=True, exist_ok=True)
    phone = {}
    for k in range(0, len(rows), 200):
        chunk = rows[k:k + 200]
        lines = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>', f'<smses count="{sum(len(c["messages"]) for c in chunk)}">']
        for c in chunk:
            authors = sorted({a for a, _ in c["messages"]})
            for a in authors:
                phone.setdefault(a, f"+9198{r.randint(10_000_000, 99_999_999)}")
            owner = authors[-1]
            c["phones"] = {a: phone[a] for a in authors}
            for (a, text), t in zip(c["messages"], timestamps(r, len(c["messages"]))):
                other = [x for x in authors if x != owner][0] if len(authors) > 1 else owner
                typ = 2 if a == owner else 1              # 1 = received, 2 = sent
                lines.append(f'  <sms protocol="0" address={quoteattr(phone[other])} date="{int(t.timestamp() * 1000)}" '
                             f'type="{typ}" body={quoteattr(text)} read="1" />')
        lines.append("</smses>")
        (out / f"sms-backup-{k // 200:03d}.xml").write_text("\n".join(lines), encoding="utf-8")
    write_jsonl(SYN / f"sms_{pool}.jsonl", rows)
    return len(rows)


# ---------------------------------------------------------------- email
GREET = ["Hi {n},", "Hello {n},", "Hey {n},", "Dear {n},"]
SIGN = ["Thanks", "Regards", "Cheers", "Take care", "Best"]


def gen_email(pool):
    rows = make(pool, N["email_threads"], 3, "email", max_beats=4)
    r = rng(33)
    out = SYN / "email" / pool
    out.mkdir(parents=True, exist_ok=True)
    buf, part = [], 0
    for idx, c in enumerate(rows):
        # merge consecutive messages from the same author into one email
        emails = []
        for a, t in c["messages"]:
            if emails and emails[-1][0] == a:
                emails[-1][1].append(t)
            else:
                emails.append([a, [t]])
        c["messages"] = [[a, " ".join(ts)] for a, ts in emails]
        addr = {a: f"{a.replace('.', '_')}@{r.choice(['gmail.com', 'yahoo.com', 'outlook.com', 'rediffmail.com'])}"
                for a, _ in emails}
        subject = r.choice(["hi", "hello", "about today", "re: class", "quick question", "hey", "this weekend"])
        for (a, body), t in zip(c["messages"], timestamps(r, len(emails))):
            other = [x for x in addr if x != a] or [a]
            m = EmailMessage()
            m["From"], m["To"] = addr[a], addr[other[0]]
            m["Subject"], m["Date"] = subject, format_datetime(t)
            m.set_content(f"{r.choice(GREET).format(n=other[0].split('_')[0])}\n\n{body}\n\n{r.choice(SIGN)},\n{a}\n")
            buf.append(f"From {addr[a]} {t.strftime('%a %b %d %H:%M:%S %Y')}\n" + m.as_string() + "\n")
        if len(buf) >= 400 or idx == len(rows) - 1:
            (out / f"threads-{part:03d}.mbox").write_text("".join(buf), encoding="utf-8")
            buf, part = [], part + 1
        c["addresses"] = addr
    write_jsonl(SYN / f"email_{pool}.jsonl", rows)
    return len(rows)


# ---------------------------------------------------------------- PDF / Word
def gen_docs(pool):
    from docx import Document
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    rows = make(pool, N["documents"], 4, "document")
    r = rng(44)
    out = SYN / "docs" / pool
    out.mkdir(parents=True, exist_ok=True)
    for c in rows:
        ts = timestamps(r, len(c["messages"]))
        lines = [f"[{t:%d/%m/%Y, %H:%M}] {a}: {m}" for (a, m), t in zip(c["messages"], ts)]
        title = r.choice(["Chat export", "Conversation transcript", "Message history", "Exported chat"])
        # PDF
        pdf = out / f"{c['id']}.pdf"
        cv = canvas.Canvas(str(pdf), pagesize=A4)
        y = 800
        cv.setFont("Helvetica-Bold", 13)
        cv.drawString(50, y, title)
        cv.setFont("Helvetica", 10)
        y -= 28
        for line in lines:
            for seg in textwrap.wrap(line, 95) or [""]:
                if y < 50:
                    cv.showPage()
                    cv.setFont("Helvetica", 10)
                    y = 800
                cv.drawString(50, y, seg)
                y -= 15
        cv.save()
        # Word
        doc = Document()
        doc.add_heading(title, level=1)
        for line in lines:
            doc.add_paragraph(line)
        doc.save(out / f"{c['id']}.docx")
        c["files"] = [f"docs/{pool}/{c['id']}.pdf", f"docs/{pool}/{c['id']}.docx"]
    write_jsonl(SYN / f"docs_{pool}.jsonl", rows)
    return len(rows)


# ---------------------------------------------------------------- screenshots
_FONTS = {}


def font(size, bold=False):
    key = (size, bold)
    if key not in _FONTS:
        _FONTS[key] = ImageFont.truetype(FONT_B if bold else FONT, size)
    return _FONTS[key]


def render_screen(messages, owner, contact, r, upto=None):
    """Phone-style chat screenshot. Owner's messages on the right (green), the
    other party on the left (white). Returns (image, messages_shown)."""
    msgs = messages[:upto] if upto else messages
    img = Image.new("RGB", (W, H), (236, 229, 221))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, 110], fill=(7, 94, 84))
    d.text((40, 38), contact, font=font(34, True), fill=(255, 255, 255))
    d.text((40, 80), "online", font=font(20), fill=(220, 240, 235))
    y, shown = 140, []
    f = font(28)
    for a, text in msgs:
        wrapped = textwrap.wrap(text, 26) or [""]
        bw = max(d.textlength(s, font=f) for s in wrapped) + 40
        bh = 38 * len(wrapped) + 30
        if y + bh > H - 40:
            break
        mine = a == owner
        x0 = W - bw - 30 if mine else 30
        d.rounded_rectangle([x0, y, x0 + bw, y + bh], radius=18, fill=(220, 248, 198) if mine else (255, 255, 255))
        for k, s in enumerate(wrapped):
            d.text((x0 + 20, y + 14 + 38 * k), s, font=f, fill=(20, 20, 20))
        d.text((x0 + bw - 70, y + bh - 22), f"{r.randint(7, 11)}:{r.randint(10, 59)}", font=font(16), fill=(120, 120, 120))
        shown.append([a, text])
        y += bh + 16
    return img, shown


def gen_screens(pool):
    rows = make(pool, N["screenshots"], 5, "screenshot", kind_mix="balanced", max_beats=6)
    r = rng(55)
    out = SYN / "screens" / pool
    out.mkdir(parents=True, exist_ok=True)
    for c in rows:
        authors = list(dict.fromkeys(a for a, _ in c["messages"]))
        owner = authors[-1] if c["label"] else r.choice(authors)       # seized phone = the child's
        other = [a for a in authors if a != owner][0] if len(authors) > 1 else owner
        img, shown = render_screen(c["messages"], owner, other, r)
        img.save(out / f"{c['id']}.jpg", quality=88)
        c.update({"messages": shown, "owner": owner, "files": [f"screens/{pool}/{c['id']}.jpg"]})
        if c["predator"] and c["predator"] not in {a for a, _ in shown}:
            c["predator"] = None
            c["label"] = 0 if not shown else c["label"]
    write_jsonl(SYN / f"screens_{pool}.jsonl", rows)
    return len(rows)


# ---------------------------------------------------------------- voice notes
# Windows built-in speech (System.Speech) via ml_v2/tts.ps1 — Piper's native library
# is blocked by Windows Smart App Control on this machine. Two system voices, each
# used at several speaking rates, give distinct-sounding speakers.
SYSTEM_VOICES = ["Microsoft David Desktop", "Microsoft Zira Desktop"]
RATES = [-2, -1, 0, 1, 2]
TTS_PS1 = Path(__file__).resolve().parent / "tts.ps1"


def speakers(r):
    """Two different (voice, rate) speakers for a conversation."""
    a = (r.choice(SYSTEM_VOICES), r.choice(RATES))
    b = (r.choice(SYSTEM_VOICES), r.choice(RATES))
    while b == a:
        b = (r.choice(SYSTEM_VOICES), r.choice(RATES))
    return a, b


def tts_batch(jobs):
    """jobs: list of (text, (voice, rate), path). Returns durations in seconds."""
    if not jobs:
        return []
    jf = jobs[0][2].parent / f"_jobs_{jobs[0][2].stem}.json"
    jf.write_text(json.dumps([{"text": t, "voice": v, "rate": rt, "path": str(p)} for t, (v, rt), p in jobs]),
                  encoding="utf-8")
    subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(TTS_PS1), str(jf)],
                   check=True, capture_output=True)
    jf.unlink()
    out = []
    for _, _, p in jobs:
        with wave.open(str(p), "rb") as wf:
            out.append(wf.getnframes() / wf.getframerate())
    return out


def gen_voice(pool):
    rows = make(pool, N["voice_convs"], 6, "voice", kind_mix="balanced", max_beats=6)
    r = rng(66)
    out = SYN / "voice" / pool
    for c in rows:
        authors = list(dict.fromkeys(a for a, _ in c["messages"]))
        sa, sb = speakers(r)
        vmap = {authors[0]: sa, **({authors[1]: sb} if len(authors) > 1 else {})}
        folder = out / c["id"]
        folder.mkdir(parents=True, exist_ok=True)
        jobs = [(text, vmap[a], folder / f"{k:02d}.wav") for k, (a, text) in enumerate(c["messages"])]
        durs = tts_batch(jobs)
        c["clips"] = [{"file": f"voice/{pool}/{c['id']}/{k:02d}.wav", "author": a, "seconds": round(d, 2)}
                      for k, ((a, _), d) in enumerate(zip(c["messages"], durs))]
        c["voices"] = {a: f"{v} (rate {rt})" for a, (v, rt) in vmap.items()}
    write_jsonl(SYN / f"voice_{pool}.jsonl", rows)
    return len(rows)


# ---------------------------------------------------------------- video
def gen_video(pool):
    import imageio_ffmpeg
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    rows = make(pool, N["videos"], 7, "video", kind_mix="balanced", max_beats=5)
    r = rng(77)
    out = SYN / "video" / pool
    work = SYN / "video" / f"_work_{pool}"
    out.mkdir(parents=True, exist_ok=True)
    work.mkdir(parents=True, exist_ok=True)
    for c in rows:
        authors = list(dict.fromkeys(a for a, _ in c["messages"]))
        owner = authors[-1]
        other = authors[0]
        sa, sb = speakers(r)
        vmap = {authors[0]: sa, **({authors[1]: sb} if len(authors) > 1 else {})}
        msgs = c["messages"][:12]
        wavs = [work / f"{c['id']}_{k:02d}.wav" for k in range(1, len(msgs) + 1)]
        durs = tts_batch([(t, vmap[a], w) for (a, t), w in zip(msgs, wavs)])
        concat, shown_all = [], []
        for k in range(1, len(msgs) + 1):
            img, shown = render_screen(msgs, owner, other, r, upto=k)
            fp = work / f"{c['id']}_{k:02d}.png"
            img.save(fp)
            concat.append(f"file '{fp.as_posix()}'\nduration {durs[k - 1] + 0.4:.2f}")
            shown_all = shown
        concat.append(f"file '{(work / f'{c['id']}_{len(msgs):02d}.png').as_posix()}'")
        lst = work / f"{c['id']}_frames.txt"
        lst.write_text("\n".join(concat), encoding="utf-8")
        alst = work / f"{c['id']}_audio.txt"
        alst.write_text("\n".join(f"file '{w.as_posix()}'" for w in wavs), encoding="utf-8")
        mp4 = out / f"{c['id']}.mp4"
        cmd = [ffmpeg, "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", str(lst),
               "-f", "concat", "-safe", "0", "-i", str(alst), "-vsync", "vfr", "-pix_fmt", "yuv420p",
               "-c:v", "libx264", "-preset", "veryfast", "-crf", "26", "-c:a", "aac", "-b:a", "96k",
               "-shortest", str(mp4)]
        subprocess.run(cmd, check=True)
        for p in work.glob(f"{c['id']}_*"):
            p.unlink()
        c.update({"messages": msgs, "visible_messages": shown_all, "owner": owner,
                  "voices": {a: f"{v} (rate {rt})" for a, (v, rt) in vmap.items()},
                  "files": [f"video/{pool}/{c['id']}.mp4"]})
    write_jsonl(SYN / f"video_{pool}.jsonl", rows)
    return len(rows)


STEPS = [("chats", gen_chats), ("sms", gen_sms), ("email", gen_email), ("docs", gen_docs),
         ("screens", gen_screens), ("voice", gen_voice), ("video", gen_video)]


def main(only=None):
    for name, fn in STEPS:
        if only and name not in only:
            continue
        for pool in ("A", "B"):
            marker = SYN / f".done_{name}_{pool}"
            if marker.exists():
                log("generate", f"skip {name} {pool} (done)")
                continue
            t0 = time.time()
            n = fn(pool)
            marker.write_text(str(n))
            log("generate", f"{name} pool {pool}: {n} records in {time.time() - t0:.0f}s")
    log("generate", "stage complete")


if __name__ == "__main__":
    import sys
    main(set(sys.argv[1:]) or None)
